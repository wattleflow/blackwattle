# Module Name: pipelines/convertors/md2word.py
# Author: (wattleflow@outlook.com)
# Copyright: © 2022–2026 WattleFlow. All rights reserved.
# License: Apache 2 Licence


"""md2word.py — Convert a Markdown file to a black-and-white Word (.docx) document.
Single-class converter. Recognised Markdown tags:
    # / ## / ###      headings (Heading 1 / 2 / 3)
    - text  or * text  bullet list item
    1. text            numbered list item
    **text**           bold (inline)
    *text*             italic (inline)
    `text`             monospace (inline)
    | a | b |          pipe table (first row = header, light-grey shading)
    --- / === / ___    horizontal rule
    blank line         paragraph / block separator
Formatting follows a strict black-and-white house style:
    body font Segoe UI 11pt, headings 16/13/12pt, all text black,
    table header cells light grey, no document metadata.
"""

# --------------------------------------------------------------------------- #
# region imports                                                              #
# --------------------------------------------------------------------------- #
from __future__ import annotations
import re
from pathlib import Path
from typing import Any, List
from wattleflow.core import IProcessor, ITarget
from wattleflow.concrete import GenericPipeline
from wattleflow.enums.event import Event
from wattleflow.documents import FileDocument

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor
except ImportError as e:
    raise RuntimeError("PipelineMarkdownToWord requires python-docx") from e
# --------------------------------------------------------------------------- #
# endregion imports                                                           #
# --------------------------------------------------------------------------- #

# --------------------------------------------------------------------------- #
# region Constants                                                            #
# --------------------------------------------------------------------------- #

DEFAULT_FONT = "Segoe UI"
DEFAULT_SIZE = 11
DEFAULT_H1 = 16
DEFAULT_H2 = 13
DEFAULT_H3 = 12
DEFAULT_AUTHOR = ""
HEADER_GREY = "D9D9D9"

# --------------------------------------------------------------------------- #
# endregion Constants                                                         #
# --------------------------------------------------------------------------- #

# --------------------------------------------------------------------------- #
# region Pipelines
# --------------------------------------------------------------------------- #


class PipelineMarkdownToWord(GenericPipeline):
    ALLOWED = ["font", "size", "h1", "h2", "h3", "author"]

    _INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
    _HR_RE = re.compile(r"^\s*([-=_])\1{2,}\s*$")
    _HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
    _BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
    _NUMBER_RE = re.compile(r"^\s*\d+\.\s+(.*)$")

    def __init__(self, **kwargs: Any) -> None:
        super().__init__(
            font=kwargs.pop("font", DEFAULT_FONT) or DEFAULT_FONT,
            size=int(kwargs.pop("size", DEFAULT_SIZE) or DEFAULT_SIZE),
            h1=int(kwargs.pop("h1", DEFAULT_H1) or DEFAULT_H1),
            h2=int(kwargs.pop("h2", DEFAULT_H2) or DEFAULT_H2),
            h3=int(kwargs.pop("h3", DEFAULT_H3) or DEFAULT_H3),
            author=kwargs.pop("author", DEFAULT_AUTHOR) or DEFAULT_AUTHOR,
            **kwargs,
        )

        self.BLACK = RGBColor(0x00, 0x00, 0x00)
        self.HEADER_GREY = HEADER_GREY
        self.doc = Document()
        self._configure_styles()

    # ---------------------------------------------------------------- #
    # Style and metadata setup                                         #
    # ---------------------------------------------------------------- #
    def _configure_styles(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(self.size)
        normal.font.color.rgb = self.BLACK

        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), self.font)

        heading_sizes = {1: self.h1, 2: self.h2, 3: self.h3}
        for level, size in heading_sizes.items():
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = self.font
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = self.BLACK

    def _strip_metadata(self) -> None:
        cp = self.doc.core_properties
        cp.author = self.author
        cp.last_modified_by = self.author
        for field in ("title", "subject", "keywords", "comments", "category"):
            try:
                setattr(cp, field, "")
            except (ValueError, AttributeError):
                pass

    # ---------------------------------------------------------------- #
    # Line-continuity helpers                                          #
    # ---------------------------------------------------------------- #
    def _is_block_start(self, line: str) -> bool:
        if line.strip() == "":
            return True
        if self._HEADING_RE.match(line) or self._HR_RE.match(line):
            return True
        if self._BULLET_RE.match(line) or self._NUMBER_RE.match(line):
            return True
        if line.lstrip().startswith("|"):
            return True
        return False

    def _collapse_list_item(self, first_text: str, continuation_lines: List[str]) -> str:
        parts = [first_text.strip()]
        parts.extend(c.strip() for c in continuation_lines)
        return re.sub(r"\s+", " ", " ".join(p for p in parts if p)).strip()

    def _collapse_paragraph(self, lines: List[str]) -> str:
        return re.sub(r"\s+", " ", " ".join(line.strip() for line in lines)).strip()

    # ----------------------------------------------------------------- #
    # Inline formatting                                                 #
    # ----------------------------------------------------------------- #
    def _add_inline(self, paragraph, text: str) -> None:
        for token in self._INLINE_RE.split(text):
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*") and token.endswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            elif token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
            else:
                run = paragraph.add_run(token)
            run.font.color.rgb = self.BLACK

    # ----------------------------------------------------------------- #
    # Block builders                                                    #
    # ----------------------------------------------------------------- #
    def add_heading(self, level: int, text: str) -> None:
        para = self.doc.add_paragraph(style=f"Heading {level}")
        self._add_inline(para, text)

    def add_paragraph(self, text: str) -> None:
        para = self.doc.add_paragraph()
        self._add_inline(para, text)

    def add_bullet(self, text: str) -> None:
        para = self.doc.add_paragraph(style="List Bullet")
        self._add_inline(para, text)

    def add_numbered(self, text: str) -> None:
        para = self.doc.add_paragraph(style="List Number")
        self._add_inline(para, text)

    def add_horizontal_rule(self) -> None:
        para = self.doc.add_paragraph()
        p_pr = para._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        borders.append(bottom)
        p_pr.append(borders)

    def add_table(self, rows: List[List[str]]) -> None:
        table = self.doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for r_idx, row in enumerate(rows):
            cells = table.add_row().cells
            for c_idx, cell_text in enumerate(row):
                cell = cells[c_idx]
                cell.text = ""
                self._add_inline(cell.paragraphs[0], cell_text.strip())
                if r_idx == 0:
                    self._shade_cell(cell, self.HEADER_GREY)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True

    def _shade_cell(self, cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    # ----------------------------------------------------------------- #
    # Main parse loop                                                   #
    # ----------------------------------------------------------------- #
    def _convert(self, input_path: str, output_path: str) -> str:
        # One pipeline instance serves every document of the cycle, so the docx
        # under construction is per-CONVERSION state, not per-instance. Built in
        # __init__ and never reset, each output carried every document processed
        # before it (13 markdowns -> a last .docx holding all 13).
        self.doc = Document()
        self._configure_styles()

        lines = Path(input_path).read_text(encoding="utf-8").splitlines()
        i, n = 0, len(lines)

        while i < n:
            line = lines[i]

            if line.strip() == "":
                i += 1
                continue

            if self._HR_RE.match(line):
                self.add_horizontal_rule()
                i += 1
                continue

            heading = self._HEADING_RE.match(line)
            if heading:
                level = len(heading.group(1))
                self.add_heading(level, heading.group(2).strip())
                i += 1
                continue

            if line.lstrip().startswith("|"):
                i = self._consume_table(lines, i)
                continue

            bullet = self._BULLET_RE.match(line)
            number = self._NUMBER_RE.match(line)
            if bullet or number:
                text, i = self._consume_list_item(lines, i, bullet or number)
                if bullet:
                    self.add_bullet(text)
                else:
                    self.add_numbered(text)
                continue

            buf = [line]
            i += 1
            while i < n and not self._is_block_start(lines[i]):
                buf.append(lines[i])
                i += 1
            self.add_paragraph(self._collapse_paragraph(buf))

        self._strip_metadata()
        self.doc.save(output_path)
        return output_path

    def _consume_list_item(self, lines: List[str], i: int, match: re.Match):
        first_text = match.group(1)
        i += 1
        continuation: List[str] = []
        n = len(lines)
        while i < n and not self._is_block_start(lines[i]):
            continuation.append(lines[i])
            i += 1
        return self._collapse_list_item(first_text, continuation), i

    def _consume_table(self, lines: List[str], i: int) -> int:
        rows: List[List[str]] = []
        n = len(lines)
        while i < n and lines[i].lstrip().startswith("|"):
            stripped = lines[i].strip()
            if re.fullmatch(r"[\s|:\-]+", stripped):
                i += 1
                continue
            raw = stripped.strip("|")
            rows.append([c.strip() for c in raw.split("|")])
            i += 1
        if rows:
            self.add_table(rows)
        return i

    # ----------------------------------------------------------------- #
    # Pipeline entry                                                    #
    # ----------------------------------------------------------------- #
    def transform(
        self,
        processor: IProcessor,
        facade: ITarget,
        *args,
        **kwargs,
    ) -> str | None:
        document: FileDocument = facade.request()
        assert isinstance(document, FileDocument), "Expected FileDocument type: Found %s" % type(
            document
        )

        filename_raw = getattr(document, "filename", None)
        filename: Path = Path(filename_raw) if filename_raw else None
        if filename is None or not filename.exists():
            self.error(
                msg=Event.Transform.name,
                step=Event.Started.name,
                reason="No valid file to process!",
                document=document,
                filename=filename,
            )
            return

        if not document.size > 0:
            self.warning(
                msg=Event.Transform.name,
                step=Event.Check.name,
                reason="Nothing to process here!",
                document=document,
                size=document.size,
            )
            return

        output = filename.with_suffix(".docx")
        self._convert(str(filename), str(output))
        document.update_metadata("output", str(output))

        uid = processor.blackboard.write(
            pipeline=self,
            facade=facade,
            processor=processor,
        )

        self.debug(
            msg=Event.Transform.name,
            step=Event.Completed.name,
            uid=uid,
            size=document.size,
            output=str(output),
        )
        # Returned, not dropped: GenericPipeline.process reports it as `result=`
        # on its DEBUG completion record; the INFO record that stands for this
        # document belongs to the processor.
        return uid


# --------------------------------------------------------------------------- #
# endregion Pipelines                                                         #
# --------------------------------------------------------------------------- #
